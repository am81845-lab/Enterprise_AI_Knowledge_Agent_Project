from __future__ import annotations
import json
import os
import re
from pathlib import Path
from typing import Iterable
import streamlit as st
from docx import Document
from pypdf import PdfReader
from langchain_text_splitters import RecursiveCharacterTextSplitter
from sentence_transformers import SentenceTransformer
import faiss
import numpy as np

ROOT = Path(__file__).resolve().parents[1]
DOCS_DIR = ROOT / "data" / "documents"
INDEX_DIR = ROOT / "data" / "index"
FAISS_PATH = INDEX_DIR / "knowledge.faiss"
CHUNKS_PATH = INDEX_DIR / "chunks.json"
MODEL_NAME = os.getenv("EMBEDDING_MODEL", "sentence-transformers/all-MiniLM-L6-v2")


def _safe_name(name: str) -> str:
    name = Path(name).name
    name = re.sub(r"[^A-Za-z0-9._ -]", "_", name).strip(" .")
    return name or "uploaded_document"


def save_uploaded_files(files: Iterable) -> list[str]:
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    saved = []
    for file in files:
        filename = _safe_name(file.name)
        target = DOCS_DIR / filename
        # Avoid silently overwriting an existing file.
        if target.exists():
            target = DOCS_DIR / f"{target.stem}_{__import__('hashlib').sha256(file.getvalue()).hexdigest()[:8]}{target.suffix}"
        target.write_bytes(file.getvalue())
        saved.append(target.name)
    return saved


def list_documents() -> list[dict]:
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    items = []
    for path in sorted(DOCS_DIR.iterdir()):
        if path.is_file() and path.suffix.lower() in {".pdf", ".docx", ".txt"}:
            items.append({"name": path.name, "size_kb": path.stat().st_size / 1024})
    return items


def extract_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".pdf":
        reader = PdfReader(str(path))
        pages = []
        for page_num, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                pages.append(f"[Page {page_num}]\n{text}")
        raw = "\n\n".join(pages)
    elif ext == ".docx":
        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip().replace("\n", " ") for cell in row.cells))
        raw = "\n".join(parts)
    elif ext == ".txt":
        raw = path.read_text(encoding="utf-8", errors="replace")
    else:
        raise ValueError(f"Unsupported file format: {ext}")
    raw = raw.replace("\x00", " ")
    raw = re.sub(r"[ \t]+", " ", raw)
    raw = re.sub(r"\n{3,}", "\n\n", raw)
    return raw.strip()


@st.cache_resource(show_spinner=False)
def get_embedding_model() -> SentenceTransformer:
    return SentenceTransformer(MODEL_NAME)


def build_index() -> int:
    INDEX_DIR.mkdir(parents=True, exist_ok=True)
    paths = [p for p in sorted(DOCS_DIR.iterdir()) if p.is_file() and p.suffix.lower() in {".pdf", ".docx", ".txt"}]
    if not paths:
        raise ValueError("Upload at least one PDF, DOCX, or TXT file before building the index.")
    splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=200, separators=["\n\n", "\n", ". ", " ", ""])
    chunks = []
    for path in paths:
        text = extract_text(path)
        if not text:
            continue
        pieces = splitter.split_text(text)
        for idx, piece in enumerate(pieces):
            chunks.append({"text": piece, "source": path.name, "chunk_id": idx + 1})
    if not chunks:
        raise ValueError("No extractable text was found. Scanned PDFs may need OCR support.")
    model = get_embedding_model()
    vectors = model.encode([c["text"] for c in chunks], normalize_embeddings=True, convert_to_numpy=True, show_progress_bar=False)
    vectors = np.asarray(vectors, dtype="float32")
    index = faiss.IndexFlatIP(vectors.shape[1])  # normalized vectors => inner product is cosine similarity
    index.add(vectors)
    faiss.write_index(index, str(FAISS_PATH))
    CHUNKS_PATH.write_text(json.dumps(chunks, ensure_ascii=False, indent=2), encoding="utf-8")
    return len(chunks)


def clear_index() -> None:
    for path in (FAISS_PATH, CHUNKS_PATH):
        if path.exists():
            path.unlink()


def load_index():
    if not FAISS_PATH.exists() or not CHUNKS_PATH.exists():
        return None, []
    index = faiss.read_index(str(FAISS_PATH))
    chunks = json.loads(CHUNKS_PATH.read_text(encoding="utf-8"))
    if index.ntotal != len(chunks):
        raise ValueError("Index metadata mismatch. Rebuild the index from the Admin Dashboard.")
    return index, chunks
